"""Vistas del módulo Presupuesto FSE (Fase 1 — Núcleo Presupuestal)."""
import logging
from decimal import Decimal, InvalidOperation

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import ProtectedError, Q
from django.http import HttpResponse, JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.views.decorators.http import require_POST

from . import reportes, services
from .forms import (
    ApropiacionForm,
    CDPForm,
    ConceptoRetencionForm,
    CuentaBancariaForm,
    ElementoAlmacenForm,
    GenerarComprobanteForm,
    ModificacionPresupuestalForm,
    MovimientoAlmacenForm,
    ObligacionForm,
    OrdenDePagoForm,
    PresupuestoIngresoForm,
    RetencionAplicadaForm,
    RPForm,
    RubroPresupuestalGastoForm,
    RubroPresupuestalIngresoForm,
    VigenciaFiscalForm,
)
from .models import (
    CDP,
    RP,
    Apropiacion,
    CatalogoGeneralCuentas,
    CategoriaCPC,
    ComprobanteContable,
    ConceptoRetencion,
    CuentaBancaria,
    ElementoAlmacen,
    ModificacionPresupuestal,
    MovimientoAlmacen,
    MovimientoTesoreria,
    Obligacion,
    OrdenDePago,
    PresupuestoIngreso,
    RetencionAplicada,
    RubroPresupuestalGasto,
    RubroPresupuestalIngreso,
    VigenciaFiscal,
)

logger = logging.getLogger(__name__)


def _get_institucion(request):
    return getattr(request.user, 'institucion_asociada', None)


def _es_gestor_presupuesto(user):
    rol = getattr(user, 'rol', '') or ''
    return rol in ('administrador', 'rector', 'tesoreria') or user.is_superuser


def _requiere_gestor(request):
    """Devuelve None si el usuario puede gestionar presupuesto, o un redirect si no."""
    if not _es_gestor_presupuesto(request.user):
        messages.error(request, 'No tienes permiso para acceder al módulo de Presupuesto.')
        return redirect('gestion_academica:inicio_academico')
    return None


def _filtro_institucion(request):
    """Igual criterio que el resto de la plataforma: superusuario ve todo
    (sin filtrar), el resto solo su propia institución."""
    if request.user.is_superuser:
        return {}
    return {'institucion': _get_institucion(request)}


def _eliminar_generico(request, model, pk, volver_url, nombre_singular):
    """Elimina un registro de catálogo simple (sin cadena presupuestal
    propia) — vigencias, rubros, conceptos de retención, cuentas bancarias,
    elementos de almacén. Si otro registro ya depende de él (protegido con
    on_delete=PROTECT), se avisa en lenguaje sencillo en vez de dejar pasar
    el error técnico de Django."""
    guard = _requiere_gestor(request)
    if guard:
        return guard
    obj = get_object_or_404(model, pk=pk, **_filtro_institucion(request))
    try:
        obj.delete()
        messages.success(request, f'{nombre_singular} eliminado.')
    except ProtectedError:
        messages.error(
            request,
            f'No se puede eliminar: este {nombre_singular.lower()} ya tiene movimientos o registros '
            'asociados. Solo se pueden eliminar los que todavía no se han usado.',
        )
    return redirect(volver_url)


# ─────────────────────────────────────────────────────────────────────────
# Dashboard
# ─────────────────────────────────────────────────────────────────────────

@login_required
def dashboard(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    filtro = _filtro_institucion(request)
    vigencia_id = request.GET.get('vigencia')
    vigencias = VigenciaFiscal.objects.filter(**filtro).order_by('-anio')
    vigencia_actual = None
    if vigencia_id:
        vigencia_actual = vigencias.filter(pk=vigencia_id).first()
    if not vigencia_actual:
        vigencia_actual = vigencias.filter(estado=VigenciaFiscal.Estado.ABIERTA).first() or vigencias.first()

    apropiaciones = []
    if vigencia_actual:
        apropiaciones = list(Apropiacion.objects.filter(vigencia=vigencia_actual).select_related('rubro'))

    total_apropiado = sum((a.valor_definitivo for a in apropiaciones), start=0)
    total_comprometido = sum((a.total_comprometido for a in apropiaciones), start=0)
    total_disponible = sum((a.saldo_disponible for a in apropiaciones), start=0)

    context = {
        'titulo_pagina': 'Presupuesto FSE',
        'vigencias': vigencias,
        'vigencia_actual': vigencia_actual,
        'apropiaciones': apropiaciones,
        'total_apropiado': total_apropiado,
        'total_comprometido': total_comprometido,
        'total_disponible': total_disponible,
    }
    return render(request, 'presupuesto/dashboard.html', context)


@login_required
def manual_uso(request):
    """Explicación en lenguaje sencillo de cada herramienta del módulo,
    agrupada en las mismas secciones que el menú lateral — para que un
    rector/tesorero no técnico entienda qué hace cada pantalla sin
    necesitar soporte técnico."""
    guard = _requiere_gestor(request)
    if guard:
        return guard
    return render(request, 'presupuesto/manual_uso.html', {
        'titulo_pagina': 'Manual de Uso — Presupuesto FSE',
    })


# ─────────────────────────────────────────────────────────────────────────
# Vigencias fiscales
# ─────────────────────────────────────────────────────────────────────────

@login_required
def lista_vigencias(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    vigencias = VigenciaFiscal.objects.filter(**_filtro_institucion(request)).order_by('-anio')
    return render(request, 'presupuesto/vigencia_lista.html', {
        'titulo_pagina': 'Vigencias Fiscales', 'vigencias': vigencias,
    })


@login_required
def crear_vigencia(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    if request.method == 'POST':
        form = VigenciaFiscalForm(request.POST, institucion=institucion)
        if form.is_valid():
            vigencia = form.save(commit=False)
            vigencia.institucion = institucion
            try:
                vigencia.full_clean()
                vigencia.save()
                messages.success(request, 'Vigencia %(anio)s abierta correctamente.' % {'anio': vigencia.anio})
                return redirect('presupuesto:lista_vigencias')
            except ValidationError as e:
                for field, errs in getattr(e, 'message_dict', {'__all__': e.messages}).items():
                    for err in errs:
                        form.add_error(field if field != '__all__' else None, err)
    else:
        form = VigenciaFiscalForm(institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Abrir Vigencia Fiscal', 'form': form,
        'icono': 'bi-calendar-range', 'volver_url': 'presupuesto:lista_vigencias',
    })


@login_required
def editar_vigencia(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    vigencia = get_object_or_404(VigenciaFiscal, pk=pk, **_filtro_institucion(request))
    if request.method == 'POST':
        form = VigenciaFiscalForm(request.POST, instance=vigencia, institucion=institucion)
        if form.is_valid():
            actualizada = form.save(commit=False)
            try:
                actualizada.full_clean()
                actualizada.save()
                messages.success(request, 'Vigencia actualizada.')
                return redirect('presupuesto:lista_vigencias')
            except ValidationError as e:
                for field, errs in getattr(e, 'message_dict', {'__all__': e.messages}).items():
                    for err in errs:
                        form.add_error(field if field != '__all__' else None, err)
    else:
        form = VigenciaFiscalForm(instance=vigencia, institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Editar Vigencia Fiscal', 'form': form,
        'icono': 'bi-calendar-range', 'volver_url': 'presupuesto:lista_vigencias',
    })


@login_required
@require_POST
def eliminar_vigencia(request, pk):
    return _eliminar_generico(request, VigenciaFiscal, pk, 'presupuesto:lista_vigencias', 'Vigencia fiscal')


@login_required
@require_POST
def cerrar_vigencia(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    filtro = _filtro_institucion(request)
    vigencia = get_object_or_404(VigenciaFiscal, pk=pk, **filtro)
    services.cerrar_vigencia(vigencia=vigencia, usuario=request.user)
    messages.success(request, 'Vigencia %(anio)s cerrada. No se podrán comprometer más recursos en ella.' % {'anio': vigencia.anio})
    return redirect('presupuesto:lista_vigencias')


# ─────────────────────────────────────────────────────────────────────────
# Rubros (ingreso / gasto)
# ─────────────────────────────────────────────────────────────────────────

@login_required
def lista_rubros_ingreso(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    rubros = RubroPresupuestalIngreso.objects.filter(**_filtro_institucion(request)).order_by('codigo')
    return render(request, 'presupuesto/rubro_ingreso_lista.html', {
        'titulo_pagina': 'Rubros de Ingreso', 'rubros': rubros,
    })


@login_required
def crear_rubro_ingreso(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    if request.method == 'POST':
        form = RubroPresupuestalIngresoForm(request.POST, institucion=institucion)
        if form.is_valid():
            rubro = form.save(commit=False)
            rubro.institucion = institucion
            rubro.save()
            messages.success(request, 'Rubro de ingreso creado.')
            return redirect('presupuesto:lista_rubros_ingreso')
    else:
        form = RubroPresupuestalIngresoForm(institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Nuevo Rubro de Ingreso', 'form': form,
        'icono': 'bi-arrow-down-circle', 'volver_url': 'presupuesto:lista_rubros_ingreso',
    })


@login_required
def editar_rubro_ingreso(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    rubro = get_object_or_404(RubroPresupuestalIngreso, pk=pk, **_filtro_institucion(request))
    if request.method == 'POST':
        form = RubroPresupuestalIngresoForm(request.POST, instance=rubro, institucion=institucion)
        if form.is_valid():
            form.save()
            messages.success(request, 'Rubro de ingreso actualizado.')
            return redirect('presupuesto:lista_rubros_ingreso')
    else:
        form = RubroPresupuestalIngresoForm(instance=rubro, institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Editar Rubro de Ingreso', 'form': form,
        'icono': 'bi-arrow-down-circle', 'volver_url': 'presupuesto:lista_rubros_ingreso',
    })


@login_required
@require_POST
def eliminar_rubro_ingreso(request, pk):
    return _eliminar_generico(request, RubroPresupuestalIngreso, pk, 'presupuesto:lista_rubros_ingreso', 'Rubro de ingreso')


@login_required
def lista_rubros_gasto(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    rubros = RubroPresupuestalGasto.objects.filter(**_filtro_institucion(request)).order_by('codigo')
    return render(request, 'presupuesto/rubro_gasto_lista.html', {
        'titulo_pagina': 'Rubros de Gasto', 'rubros': rubros,
    })


@login_required
def crear_rubro_gasto(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    if request.method == 'POST':
        form = RubroPresupuestalGastoForm(request.POST, institucion=institucion)
        if form.is_valid():
            rubro = form.save(commit=False)
            rubro.institucion = institucion
            rubro.save()
            messages.success(request, 'Rubro de gasto creado.')
            return redirect('presupuesto:lista_rubros_gasto')
    else:
        form = RubroPresupuestalGastoForm(institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Nuevo Rubro de Gasto', 'form': form,
        'icono': 'bi-arrow-up-circle', 'volver_url': 'presupuesto:lista_rubros_gasto',
    })


@login_required
def editar_rubro_gasto(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    rubro = get_object_or_404(RubroPresupuestalGasto, pk=pk, **_filtro_institucion(request))
    if request.method == 'POST':
        form = RubroPresupuestalGastoForm(request.POST, instance=rubro, institucion=institucion)
        if form.is_valid():
            form.save()
            messages.success(request, 'Rubro de gasto actualizado.')
            return redirect('presupuesto:lista_rubros_gasto')
    else:
        form = RubroPresupuestalGastoForm(instance=rubro, institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Editar Rubro de Gasto', 'form': form,
        'icono': 'bi-arrow-up-circle', 'volver_url': 'presupuesto:lista_rubros_gasto',
    })


@login_required
@require_POST
def eliminar_rubro_gasto(request, pk):
    return _eliminar_generico(request, RubroPresupuestalGasto, pk, 'presupuesto:lista_rubros_gasto', 'Rubro de gasto')


# ─────────────────────────────────────────────────────────────────────────
# Presupuesto de ingresos
# ─────────────────────────────────────────────────────────────────────────

@login_required
def lista_presupuesto_ingreso(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    items = PresupuestoIngreso.objects.filter(**_filtro_institucion(request)).select_related('rubro', 'vigencia').order_by('-vigencia__anio')
    return render(request, 'presupuesto/presupuesto_ingreso_lista.html', {
        'titulo_pagina': 'Presupuesto de Ingresos', 'items': items,
    })


@login_required
def crear_presupuesto_ingreso(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    if request.method == 'POST':
        form = PresupuestoIngresoForm(request.POST, institucion=institucion)
        if form.is_valid():
            item = form.save(commit=False)
            item.institucion = institucion
            try:
                item.full_clean()
                item.save()
                messages.success(request, 'Presupuesto de ingreso registrado.')
                return redirect('presupuesto:lista_presupuesto_ingreso')
            except ValidationError as e:
                for field, errs in getattr(e, 'message_dict', {'__all__': e.messages}).items():
                    for err in errs:
                        form.add_error(field if field != '__all__' else None, err)
    else:
        form = PresupuestoIngresoForm(institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Registrar Presupuesto de Ingreso', 'form': form,
        'icono': 'bi-cash-coin', 'volver_url': 'presupuesto:lista_presupuesto_ingreso',
    })


@login_required
def editar_presupuesto_ingreso(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    item = get_object_or_404(PresupuestoIngreso, pk=pk, **_filtro_institucion(request))
    if request.method == 'POST':
        form = PresupuestoIngresoForm(request.POST, instance=item, institucion=institucion)
        if form.is_valid():
            actualizado = form.save(commit=False)
            try:
                actualizado.full_clean()
                actualizado.save()
                messages.success(request, 'Presupuesto de ingreso actualizado.')
                return redirect('presupuesto:lista_presupuesto_ingreso')
            except ValidationError as e:
                for field, errs in getattr(e, 'message_dict', {'__all__': e.messages}).items():
                    for err in errs:
                        form.add_error(field if field != '__all__' else None, err)
    else:
        form = PresupuestoIngresoForm(instance=item, institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Editar Presupuesto de Ingreso', 'form': form,
        'icono': 'bi-cash-coin', 'volver_url': 'presupuesto:lista_presupuesto_ingreso',
    })


@login_required
@require_POST
def eliminar_presupuesto_ingreso(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    item = get_object_or_404(PresupuestoIngreso, pk=pk, **_filtro_institucion(request))
    if item.valor_recaudado and item.valor_recaudado > 0:
        messages.error(
            request,
            'No se puede eliminar: este presupuesto de ingreso ya tiene recaudo registrado '
            '(${:,.2f}). Corrígelo editándolo en vez de eliminarlo.'.format(item.valor_recaudado),
        )
        return redirect('presupuesto:lista_presupuesto_ingreso')
    item.delete()
    messages.success(request, 'Presupuesto de ingreso eliminado.')
    return redirect('presupuesto:lista_presupuesto_ingreso')


@login_required
@require_POST
def registrar_recaudo(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    item = get_object_or_404(PresupuestoIngreso, pk=pk, **_filtro_institucion(request))
    try:
        valor = Decimal(request.POST.get('valor', '0'))
        services.registrar_recaudo(presupuesto_ingreso=item, valor=valor, usuario=request.user)
        messages.success(request, 'Recaudo registrado.')
    except (InvalidOperation, ValidationError) as e:
        mensaje = e.message if hasattr(e, 'message') else str(e)
        messages.error(request, f'No se pudo registrar el recaudo: {mensaje}')
    return redirect('presupuesto:lista_presupuesto_ingreso')


# ─────────────────────────────────────────────────────────────────────────
# Apropiaciones y modificaciones
# ─────────────────────────────────────────────────────────────────────────

@login_required
def lista_apropiaciones(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    apropiaciones = Apropiacion.objects.filter(**_filtro_institucion(request)).select_related('rubro', 'vigencia').order_by('-vigencia__anio', 'rubro__codigo')
    return render(request, 'presupuesto/apropiacion_lista.html', {
        'titulo_pagina': 'Apropiaciones Presupuestales', 'apropiaciones': apropiaciones,
    })


@login_required
def crear_apropiacion(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    if request.method == 'POST':
        form = ApropiacionForm(request.POST, institucion=institucion)
        if form.is_valid():
            apropiacion = form.save(commit=False)
            apropiacion.institucion = institucion
            try:
                apropiacion.full_clean()
                apropiacion.save()
                messages.success(request, 'Apropiación registrada.')
                return redirect('presupuesto:lista_apropiaciones')
            except ValidationError as e:
                for field, errs in getattr(e, 'message_dict', {'__all__': e.messages}).items():
                    for err in errs:
                        form.add_error(field if field != '__all__' else None, err)
    else:
        form = ApropiacionForm(institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Nueva Apropiación', 'form': form,
        'icono': 'bi-piggy-bank', 'volver_url': 'presupuesto:lista_apropiaciones',
    })


@login_required
def lista_modificaciones(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    mods = ModificacionPresupuestal.objects.filter(**_filtro_institucion(request)).select_related('apropiacion', 'apropiacion_destino').order_by('-fecha')
    return render(request, 'presupuesto/modificacion_lista.html', {
        'titulo_pagina': 'Modificaciones Presupuestales', 'modificaciones': mods,
    })


@login_required
def crear_modificacion(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    if request.method == 'POST':
        form = ModificacionPresupuestalForm(request.POST, request.FILES, institucion=institucion)
        if form.is_valid():
            mod = form.save(commit=False)
            mod.institucion = institucion
            mod.creado_por = request.user
            try:
                mod.full_clean()
                mod.save()
                messages.success(request, 'Modificación presupuestal registrada.')
                return redirect('presupuesto:lista_modificaciones')
            except ValidationError as e:
                for field, errs in getattr(e, 'message_dict', {'__all__': e.messages}).items():
                    for err in errs:
                        form.add_error(field if field != '__all__' else None, err)
    else:
        form = ModificacionPresupuestalForm(institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Nueva Modificación Presupuestal', 'form': form,
        'icono': 'bi-arrow-left-right', 'volver_url': 'presupuesto:lista_modificaciones',
    })


def _validar_apropiaciones_saldo_no_negativo(*apropiaciones):
    """Después de editar/eliminar una modificación presupuestal, el saldo
    disponible de cada apropiación afectada (origen y, si es traslado,
    destino) se recalcula en vivo — si ya hay CDP expedidos contra el valor
    que se está quitando, el saldo quedaría negativo. Se bloquea el cambio
    en ese caso en vez de dejar una apropiación con saldo imposible."""
    for aprop in apropiaciones:
        if aprop is not None and aprop.saldo_disponible < 0:
            raise ValidationError(
                'No se puede aplicar este cambio: dejaría el saldo disponible de "%(rubro)s" en negativo, '
                'porque ya hay CDP expedidos contra el valor que se está quitando o cambiando. '
                'Anula primero esos CDP y vuelve a intentarlo.' % {'rubro': aprop.rubro}
            )


@login_required
def editar_modificacion(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    mod = get_object_or_404(ModificacionPresupuestal, pk=pk, **_filtro_institucion(request))
    apropiaciones_originales = [mod.apropiacion, mod.apropiacion_destino]
    if request.method == 'POST':
        form = ModificacionPresupuestalForm(request.POST, request.FILES, instance=mod, institucion=institucion)
        if form.is_valid():
            try:
                with transaction.atomic():
                    actualizada = form.save(commit=False)
                    actualizada.full_clean()
                    actualizada.save()
                    _validar_apropiaciones_saldo_no_negativo(
                        actualizada.apropiacion, actualizada.apropiacion_destino, *apropiaciones_originales
                    )
                messages.success(request, 'Modificación presupuestal actualizada.')
                return redirect('presupuesto:lista_modificaciones')
            except ValidationError as e:
                for field, errs in getattr(e, 'message_dict', {'__all__': e.messages}).items():
                    for err in errs:
                        form.add_error(field if field != '__all__' else None, err)
    else:
        form = ModificacionPresupuestalForm(instance=mod, institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Editar Modificación Presupuestal', 'form': form,
        'icono': 'bi-arrow-left-right', 'volver_url': 'presupuesto:lista_modificaciones',
    })


@login_required
@require_POST
def eliminar_modificacion(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    mod = get_object_or_404(ModificacionPresupuestal, pk=pk, **_filtro_institucion(request))
    apropiacion, apropiacion_destino = mod.apropiacion, mod.apropiacion_destino
    try:
        with transaction.atomic():
            mod.delete()
            _validar_apropiaciones_saldo_no_negativo(apropiacion, apropiacion_destino)
        messages.success(request, 'Modificación presupuestal eliminada.')
    except ValidationError as e:
        messages.error(request, e.messages[0] if getattr(e, 'messages', None) else str(e))
    return redirect('presupuesto:lista_modificaciones')


# ─────────────────────────────────────────────────────────────────────────
# CDP
# ─────────────────────────────────────────────────────────────────────────

@login_required
def lista_cdp(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    cdps = CDP.objects.filter(**_filtro_institucion(request)).select_related('apropiacion', 'vigencia').order_by('-numero')
    return render(request, 'presupuesto/cdp_lista.html', {'titulo_pagina': 'Certificados de Disponibilidad (CDP)', 'cdps': cdps})


@login_required
def crear_cdp(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    if request.method == 'POST':
        form = CDPForm(request.POST, institucion=institucion)
        if form.is_valid():
            try:
                services.expedir_cdp(
                    apropiacion=form.cleaned_data['apropiacion'],
                    valor=form.cleaned_data['valor'],
                    objeto=form.cleaned_data['objeto'],
                    usuario=request.user,
                )
                messages.success(request, 'CDP expedido correctamente.')
                return redirect('presupuesto:lista_cdp')
            except ValidationError as e:
                form.add_error(None, e.message if hasattr(e, 'message') else str(e))
    else:
        form = CDPForm(institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Expedir CDP', 'form': form,
        'icono': 'bi-file-earmark-lock', 'volver_url': 'presupuesto:lista_cdp',
    })


@login_required
@require_POST
def anular_cdp(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    cdp = get_object_or_404(CDP, pk=pk, **_filtro_institucion(request))
    if cdp.total_comprometido_rp > 0:
        messages.error(request, 'No se puede anular: ya tiene Registros Presupuestales (RP) asociados.')
    else:
        cdp.estado = CDP.Estado.ANULADO
        cdp.anulado_motivo = request.POST.get('motivo', '')
        cdp.save(update_fields=['estado', 'anulado_motivo'])
        messages.success(request, 'CDP #%(num)s anulado.' % {'num': cdp.numero})
    return redirect('presupuesto:lista_cdp')


# ─────────────────────────────────────────────────────────────────────────
# RP
# ─────────────────────────────────────────────────────────────────────────

@login_required
def lista_rp(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    rps = RP.objects.filter(**_filtro_institucion(request)).select_related('cdp', 'tercero', 'categoria_cpc').order_by('-numero')
    return render(request, 'presupuesto/rp_lista.html', {'titulo_pagina': 'Registros Presupuestales (RP)', 'rps': rps})


@login_required
def crear_rp(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    if request.method == 'POST':
        form = RPForm(request.POST, institucion=institucion)
        if form.is_valid():
            try:
                services.crear_rp(
                    cdp=form.cleaned_data['cdp'],
                    tercero=form.cleaned_data['tercero'],
                    objeto_contrato=form.cleaned_data['objeto_contrato'],
                    categoria_cpc=form.cleaned_data.get('categoria_cpc'),
                    valor=form.cleaned_data['valor'],
                    usuario=request.user,
                )
                messages.success(request, 'Registro Presupuestal (RP) creado correctamente.')
                return redirect('presupuesto:lista_rp')
            except ValidationError as e:
                form.add_error(None, e.message if hasattr(e, 'message') else str(e))
    else:
        form = RPForm(institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Nuevo Registro Presupuestal (RP)', 'form': form,
        'icono': 'bi-file-earmark-text', 'volver_url': 'presupuesto:lista_rp',
    })


@login_required
@require_POST
def anular_rp(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    rp = get_object_or_404(RP, pk=pk, **_filtro_institucion(request))
    if rp.total_obligado > 0:
        messages.error(request, 'No se puede anular: ya tiene Obligaciones asociadas.')
    else:
        rp.estado = RP.Estado.ANULADO
        rp.anulado_motivo = request.POST.get('motivo', '')
        rp.save(update_fields=['estado', 'anulado_motivo'])
        if rp.cdp.estado == CDP.Estado.AGOTADO:
            rp.cdp.estado = CDP.Estado.VIGENTE
            rp.cdp.save(update_fields=['estado'])
        messages.success(request, 'RP #%(num)s anulado.' % {'num': rp.numero})
    return redirect('presupuesto:lista_rp')


# ─────────────────────────────────────────────────────────────────────────
# Obligaciones
# ─────────────────────────────────────────────────────────────────────────

@login_required
def lista_obligaciones(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    obligaciones = Obligacion.objects.filter(**_filtro_institucion(request)).select_related('rp', 'rp__tercero').order_by('-numero')
    return render(request, 'presupuesto/obligacion_lista.html', {'titulo_pagina': 'Obligaciones', 'obligaciones': obligaciones})


@login_required
def crear_obligacion(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    if request.method == 'POST':
        form = ObligacionForm(request.POST, request.FILES, institucion=institucion)
        if form.is_valid():
            try:
                services.causar_obligacion(
                    rp=form.cleaned_data['rp'],
                    valor=form.cleaned_data['valor'],
                    soporte=form.cleaned_data.get('soporte_recibido_satisfaccion'),
                    usuario=request.user,
                )
                messages.success(request, 'Obligación causada correctamente.')
                return redirect('presupuesto:lista_obligaciones')
            except ValidationError as e:
                form.add_error(None, e.message if hasattr(e, 'message') else str(e))
    else:
        form = ObligacionForm(institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Causar Obligación', 'form': form,
        'icono': 'bi-file-earmark-check', 'volver_url': 'presupuesto:lista_obligaciones',
    })


@login_required
@require_POST
def anular_obligacion(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    obligacion = get_object_or_404(Obligacion, pk=pk, **_filtro_institucion(request))
    if obligacion.total_pagado > 0:
        messages.error(request, 'No se puede anular: ya tiene Órdenes de Pago asociadas.')
    else:
        obligacion.estado = Obligacion.Estado.ANULADA
        obligacion.anulado_motivo = request.POST.get('motivo', '')
        obligacion.save(update_fields=['estado', 'anulado_motivo'])
        if obligacion.rp.estado == RP.Estado.LIQUIDADO:
            obligacion.rp.estado = RP.Estado.VIGENTE
            obligacion.rp.save(update_fields=['estado'])
        messages.success(request, 'Obligación #%(num)s anulada.' % {'num': obligacion.numero})
    return redirect('presupuesto:lista_obligaciones')


# ─────────────────────────────────────────────────────────────────────────
# Órdenes de Pago
# ─────────────────────────────────────────────────────────────────────────

@login_required
def lista_ordenes_pago(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    ordenes = OrdenDePago.objects.filter(**_filtro_institucion(request)).select_related('obligacion', 'beneficiario', 'comprobante_contable').order_by('-numero')
    return render(request, 'presupuesto/orden_pago_lista.html', {'titulo_pagina': 'Órdenes de Pago', 'ordenes': ordenes})


@login_required
def crear_orden_pago(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    if request.method == 'POST':
        form = OrdenDePagoForm(request.POST, institucion=institucion)
        if form.is_valid():
            try:
                orden = services.generar_orden_pago(
                    obligacion=form.cleaned_data['obligacion'],
                    usuario=request.user,
                )
                messages.success(request, 'Orden de Pago generada. Ahora puedes agregar retenciones si aplican y generar el comprobante contable.')
                return redirect('presupuesto:detalle_orden_pago', pk=orden.pk)
            except ValidationError as e:
                form.add_error(None, e.message if hasattr(e, 'message') else str(e))
    else:
        form = OrdenDePagoForm(institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Generar Orden de Pago', 'form': form,
        'icono': 'bi-cash-stack', 'volver_url': 'presupuesto:lista_ordenes_pago',
    })


@login_required
def detalle_orden_pago(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    orden = get_object_or_404(
        OrdenDePago.objects.select_related('obligacion', 'beneficiario', 'comprobante_contable'),
        pk=pk, **_filtro_institucion(request),
    )
    institucion = _get_institucion(request) if not request.user.is_superuser else orden.institucion
    retenciones = orden.retenciones.select_related('concepto')
    form_retencion = RetencionAplicadaForm(institucion=institucion)
    tiene_comprobante = hasattr(orden, 'comprobante_contable')
    return render(request, 'presupuesto/orden_pago_detalle.html', {
        'titulo_pagina': f'Orden de Pago #{orden.numero}', 'orden': orden,
        'retenciones': retenciones, 'form_retencion': form_retencion,
        'tiene_comprobante': tiene_comprobante,
    })


@login_required
@require_POST
def agregar_retencion(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    orden = get_object_or_404(OrdenDePago, pk=pk, **_filtro_institucion(request))
    institucion = _get_institucion(request) if not request.user.is_superuser else orden.institucion
    form = RetencionAplicadaForm(request.POST, institucion=institucion)
    if form.is_valid():
        try:
            services.agregar_retencion(
                orden_pago=orden, concepto=form.cleaned_data['concepto'],
                base_gravable=form.cleaned_data['base_gravable'], usuario=request.user,
            )
            messages.success(request, 'Retención agregada.')
        except ValidationError as e:
            messages.error(request, e.message if hasattr(e, 'message') else str(e))
    else:
        messages.error(request, 'Datos inválidos para la retención.')
    return redirect('presupuesto:detalle_orden_pago', pk=pk)


@login_required
@require_POST
def quitar_retencion(request, pk, retencion_pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    orden = get_object_or_404(OrdenDePago, pk=pk, **_filtro_institucion(request))
    retencion = get_object_or_404(RetencionAplicada, pk=retencion_pk, orden_pago=orden)
    try:
        services.quitar_retencion(retencion=retencion, usuario=request.user)
        messages.success(request, 'Retención eliminada.')
    except ValidationError as e:
        messages.error(request, e.message if hasattr(e, 'message') else str(e))
    return redirect('presupuesto:detalle_orden_pago', pk=pk)


@login_required
def generar_comprobante(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    orden = get_object_or_404(OrdenDePago, pk=pk, **_filtro_institucion(request))
    institucion = _get_institucion(request)
    if request.method == 'POST':
        form = GenerarComprobanteForm(request.POST, institucion=institucion)
        if form.is_valid():
            try:
                comprobante = services.generar_comprobante_contable(
                    orden_pago=orden, cuenta_bancaria=form.cleaned_data['cuenta_bancaria'], usuario=request.user,
                )
                messages.success(request, 'Comprobante contable generado en Borrador. Revísalo y contabilízalo.')
                return redirect('presupuesto:detalle_comprobante', pk=comprobante.pk)
            except ValidationError as e:
                form.add_error(None, e.message if hasattr(e, 'message') else str(e))
    else:
        form = GenerarComprobanteForm(institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': f'Generar Comprobante — Orden de Pago #{orden.numero}', 'form': form,
        'icono': 'bi-journal-text', 'volver_href': reverse('presupuesto:detalle_orden_pago', args=[orden.pk]),
    })


@login_required
@require_POST
def anular_orden_pago(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    orden = get_object_or_404(OrdenDePago, pk=pk, **_filtro_institucion(request))
    if hasattr(orden, 'comprobante_contable'):
        messages.error(request, 'No se puede anular: ya tiene un comprobante contable. Usa la opción de reversar el comprobante en su lugar.')
        return redirect('presupuesto:lista_ordenes_pago')
    orden.estado = OrdenDePago.Estado.ANULADA
    orden.anulado_motivo = request.POST.get('motivo', '')
    orden.save(update_fields=['estado', 'anulado_motivo'])
    if orden.obligacion.estado == Obligacion.Estado.PAGADA:
        orden.obligacion.estado = Obligacion.Estado.VIGENTE
        orden.obligacion.save(update_fields=['estado'])
    messages.success(request, 'Orden de Pago #%(num)s anulada.' % {'num': orden.numero})
    return redirect('presupuesto:lista_ordenes_pago')


# ─────────────────────────────────────────────────────────────────────────
# Conceptos de Retención
# ─────────────────────────────────────────────────────────────────────────

@login_required
def lista_conceptos_retencion(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    conceptos = ConceptoRetencion.objects.filter(**_filtro_institucion(request)).select_related('cuenta_puc_pasivo')
    return render(request, 'presupuesto/concepto_retencion_lista.html', {
        'titulo_pagina': 'Conceptos de Retención', 'conceptos': conceptos,
    })


@login_required
def crear_concepto_retencion(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    if request.method == 'POST':
        form = ConceptoRetencionForm(request.POST, institucion=institucion)
        if form.is_valid():
            concepto = form.save(commit=False)
            concepto.institucion = institucion
            concepto.save()
            messages.success(request, 'Concepto de retención creado.')
            return redirect('presupuesto:lista_conceptos_retencion')
    else:
        form = ConceptoRetencionForm(institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Nuevo Concepto de Retención', 'form': form,
        'icono': 'bi-percent', 'volver_url': 'presupuesto:lista_conceptos_retencion',
    })


@login_required
def editar_concepto_retencion(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    concepto = get_object_or_404(ConceptoRetencion, pk=pk, **_filtro_institucion(request))
    if request.method == 'POST':
        form = ConceptoRetencionForm(request.POST, instance=concepto, institucion=institucion)
        if form.is_valid():
            form.save()
            messages.success(request, 'Concepto de retención actualizado.')
            return redirect('presupuesto:lista_conceptos_retencion')
    else:
        form = ConceptoRetencionForm(instance=concepto, institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Editar Concepto de Retención', 'form': form,
        'icono': 'bi-percent', 'volver_url': 'presupuesto:lista_conceptos_retencion',
    })


@login_required
@require_POST
def eliminar_concepto_retencion(request, pk):
    return _eliminar_generico(request, ConceptoRetencion, pk, 'presupuesto:lista_conceptos_retencion', 'Concepto de retención')


# ─────────────────────────────────────────────────────────────────────────
# Comprobantes Contables
# ─────────────────────────────────────────────────────────────────────────

@login_required
def lista_comprobantes(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    comprobantes = ComprobanteContable.objects.filter(**_filtro_institucion(request)).select_related('orden_pago', 'vigencia').order_by('-numero')
    return render(request, 'presupuesto/comprobante_lista.html', {
        'titulo_pagina': 'Comprobantes Contables', 'comprobantes': comprobantes,
    })


@login_required
def detalle_comprobante(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    comprobante = get_object_or_404(
        ComprobanteContable.objects.select_related('orden_pago', 'comprobante_que_reversa'),
        pk=pk, **_filtro_institucion(request),
    )
    movimientos = comprobante.movimientos.select_related('cuenta', 'tercero')
    return render(request, 'presupuesto/comprobante_detalle.html', {
        'titulo_pagina': f'Comprobante #{comprobante.numero}', 'comprobante': comprobante, 'movimientos': movimientos,
    })


@login_required
@require_POST
def contabilizar_comprobante(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    comprobante = get_object_or_404(ComprobanteContable, pk=pk, **_filtro_institucion(request))
    try:
        services.contabilizar_comprobante(comprobante=comprobante, usuario=request.user)
        messages.success(request, 'Comprobante contabilizado. Ya no se puede editar — a partir de ahora solo se reversa con un ajuste.')
    except ValidationError as e:
        messages.error(request, e.message if hasattr(e, 'message') else str(e))
    return redirect('presupuesto:detalle_comprobante', pk=pk)


@login_required
@require_POST
def reversar_comprobante(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    comprobante = get_object_or_404(ComprobanteContable, pk=pk, **_filtro_institucion(request))
    motivo = request.POST.get('motivo', '')
    try:
        reverso = services.reversar_comprobante(comprobante=comprobante, usuario=request.user, motivo=motivo)
        messages.success(request, f'Comprobante reversado con el ajuste #{reverso.numero}.')
        return redirect('presupuesto:detalle_comprobante', pk=reverso.pk)
    except ValidationError as e:
        messages.error(request, e.message if hasattr(e, 'message') else str(e))
        return redirect('presupuesto:detalle_comprobante', pk=pk)


# ─────────────────────────────────────────────────────────────────────────
# Catálogo General de Cuentas (solo lectura — lo administra el propietario)
# ─────────────────────────────────────────────────────────────────────────

@login_required
def lista_catalogo_cgc(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    cuentas = CatalogoGeneralCuentas.objects.filter(activo=True).select_related('cuenta_padre')
    return render(request, 'presupuesto/catalogo_cgc_lista.html', {
        'titulo_pagina': 'Catálogo General de Cuentas (CGC)', 'cuentas': cuentas,
    })


# ─────────────────────────────────────────────────────────────────────────
# Categoría CPC (DANE) — catálogo GLOBAL de ~9.933 códigos, solo lectura,
# consumido por el buscador de RPForm (ver presupuesto/widgets.py).
# ─────────────────────────────────────────────────────────────────────────

@login_required
def buscar_categoria_cpc(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    q = (request.GET.get('q') or '').strip()
    if len(q) < 2:
        return JsonResponse({'resultados': []})
    coincidencias = CategoriaCPC.objects.filter(
        Q(codigo__icontains=q) | Q(titulo__icontains=q)
    ).order_by('codigo')[:20]
    return JsonResponse({'resultados': [
        {'id': c.pk, 'texto': f"{c.codigo} · {c.titulo}"} for c in coincidencias
    ]})


# ─────────────────────────────────────────────────────────────────────────
# Fase 3 — Tesorería
# ─────────────────────────────────────────────────────────────────────────

@login_required
def lista_cuentas_bancarias(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    cuentas = CuentaBancaria.objects.filter(**_filtro_institucion(request)).select_related('cuenta_cgc')
    return render(request, 'presupuesto/cuenta_bancaria_lista.html', {
        'titulo_pagina': 'Cuentas Bancarias', 'cuentas': cuentas,
    })


@login_required
def crear_cuenta_bancaria(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    if request.method == 'POST':
        form = CuentaBancariaForm(request.POST, institucion=institucion)
        if form.is_valid():
            cuenta = form.save(commit=False)
            cuenta.institucion = institucion
            cuenta.save()
            messages.success(request, 'Cuenta bancaria creada.')
            return redirect('presupuesto:lista_cuentas_bancarias')
    else:
        form = CuentaBancariaForm(institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Nueva Cuenta Bancaria', 'form': form,
        'icono': 'bi-bank', 'volver_url': 'presupuesto:lista_cuentas_bancarias',
    })


@login_required
def editar_cuenta_bancaria(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    cuenta = get_object_or_404(CuentaBancaria, pk=pk, **_filtro_institucion(request))
    if request.method == 'POST':
        form = CuentaBancariaForm(request.POST, instance=cuenta, institucion=institucion)
        if form.is_valid():
            form.save()
            messages.success(request, 'Cuenta bancaria actualizada.')
            return redirect('presupuesto:lista_cuentas_bancarias')
    else:
        form = CuentaBancariaForm(instance=cuenta, institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Editar Cuenta Bancaria', 'form': form,
        'icono': 'bi-bank', 'volver_url': 'presupuesto:lista_cuentas_bancarias',
    })


@login_required
@require_POST
def eliminar_cuenta_bancaria(request, pk):
    return _eliminar_generico(request, CuentaBancaria, pk, 'presupuesto:lista_cuentas_bancarias', 'Cuenta bancaria')


@login_required
def lista_movimientos_tesoreria(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    movimientos = MovimientoTesoreria.objects.filter(
        **_filtro_institucion(request)
    ).select_related('cuenta_bancaria', 'comprobante_contable')
    cuenta_id = request.GET.get('cuenta')
    if cuenta_id:
        movimientos = movimientos.filter(cuenta_bancaria_id=cuenta_id)
    cuentas = CuentaBancaria.objects.filter(**_filtro_institucion(request))
    return render(request, 'presupuesto/movimiento_tesoreria_lista.html', {
        'titulo_pagina': 'Movimientos de Tesorería', 'movimientos': movimientos,
        'cuentas': cuentas, 'cuenta_seleccionada': cuenta_id,
    })


@login_required
@require_POST
def conciliar_movimiento_tesoreria(request, pk):
    from django.utils import timezone
    guard = _requiere_gestor(request)
    if guard:
        return guard
    movimiento = get_object_or_404(MovimientoTesoreria, pk=pk, **_filtro_institucion(request))
    movimiento.conciliado = True
    movimiento.fecha_conciliacion = timezone.now()
    movimiento.save(update_fields=['conciliado', 'fecha_conciliacion'])
    messages.success(request, 'Movimiento marcado como conciliado.')
    return redirect('presupuesto:lista_movimientos_tesoreria')


# ─────────────────────────────────────────────────────────────────────────
# Fase 3 — Almacén
# ─────────────────────────────────────────────────────────────────────────

@login_required
def lista_elementos_almacen(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    elementos = ElementoAlmacen.objects.filter(**_filtro_institucion(request))
    return render(request, 'presupuesto/elemento_almacen_lista.html', {
        'titulo_pagina': 'Elementos de Almacén', 'elementos': elementos,
    })


@login_required
def crear_elemento_almacen(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    if request.method == 'POST':
        form = ElementoAlmacenForm(request.POST, institucion=institucion)
        if form.is_valid():
            elemento = form.save(commit=False)
            elemento.institucion = institucion
            elemento.save()
            messages.success(request, 'Elemento de almacén creado.')
            return redirect('presupuesto:lista_elementos_almacen')
    else:
        form = ElementoAlmacenForm(institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Nuevo Elemento de Almacén', 'form': form,
        'icono': 'bi-box-seam', 'volver_url': 'presupuesto:lista_elementos_almacen',
    })


@login_required
def editar_elemento_almacen(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    elemento = get_object_or_404(ElementoAlmacen, pk=pk, **_filtro_institucion(request))
    if request.method == 'POST':
        form = ElementoAlmacenForm(request.POST, instance=elemento, institucion=institucion)
        if form.is_valid():
            form.save()
            messages.success(request, 'Elemento de almacén actualizado.')
            return redirect('presupuesto:lista_elementos_almacen')
    else:
        form = ElementoAlmacenForm(instance=elemento, institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Editar Elemento de Almacén', 'form': form,
        'icono': 'bi-box-seam', 'volver_url': 'presupuesto:lista_elementos_almacen',
    })


@login_required
@require_POST
def eliminar_elemento_almacen(request, pk):
    return _eliminar_generico(request, ElementoAlmacen, pk, 'presupuesto:lista_elementos_almacen', 'Elemento de almacén')


@login_required
def lista_movimientos_almacen(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    movimientos = MovimientoAlmacen.objects.filter(
        **_filtro_institucion(request)
    ).select_related('elemento', 'rp').order_by('-fecha')
    return render(request, 'presupuesto/movimiento_almacen_lista.html', {
        'titulo_pagina': 'Movimientos de Almacén', 'movimientos': movimientos,
    })


@login_required
def crear_movimiento_almacen(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    institucion = _get_institucion(request)
    if request.method == 'POST':
        form = MovimientoAlmacenForm(request.POST, institucion=institucion)
        if form.is_valid():
            try:
                services.registrar_movimiento_almacen(
                    elemento=form.cleaned_data['elemento'], tipo=form.cleaned_data['tipo'],
                    cantidad=form.cleaned_data['cantidad'], valor_unitario=form.cleaned_data['valor_unitario'],
                    rp=form.cleaned_data['rp'], responsable=form.cleaned_data['responsable'], usuario=request.user,
                )
                messages.success(request, 'Movimiento de almacén registrado.')
                return redirect('presupuesto:lista_movimientos_almacen')
            except ValidationError as e:
                form.add_error(None, e.message if hasattr(e, 'message') else str(e))
    else:
        form = MovimientoAlmacenForm(institucion=institucion)
    return render(request, 'presupuesto/form_generico.html', {
        'titulo_pagina': 'Registrar Movimiento de Almacén', 'form': form,
        'icono': 'bi-box-arrow-in-down', 'volver_url': 'presupuesto:lista_movimientos_almacen',
    })


# ─────────────────────────────────────────────────────────────────────────
# Fase 4 — Reportes de ejecución presupuestal y consolidación (CHIP/SIA)
# ─────────────────────────────────────────────────────────────────────────

def _resolver_vigencia(request, filtro):
    vigencia_id = request.GET.get('vigencia')
    vigencias = VigenciaFiscal.objects.filter(**filtro).order_by('-anio')
    vigencia = None
    if vigencia_id:
        vigencia = vigencias.filter(pk=vigencia_id).first()
    if not vigencia:
        vigencia = vigencias.filter(estado=VigenciaFiscal.Estado.ABIERTA).first() or vigencias.first()
    return vigencia, vigencias


@login_required
def reporte_ejecucion(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    filtro = _filtro_institucion(request)
    vigencia, vigencias = _resolver_vigencia(request, filtro)

    filas_ingresos = reportes.ejecucion_ingresos(vigencia) if vigencia else []
    filas_gastos = reportes.ejecucion_gastos(vigencia) if vigencia else []
    filas_balance = reportes.balance_comprobacion(vigencia) if vigencia else []

    return render(request, 'presupuesto/reporte_ejecucion.html', {
        'titulo_pagina': 'Reportes de Ejecución (CHIP/SIA)',
        'vigencias': vigencias, 'vigencia_actual': vigencia,
        'filas_ingresos': filas_ingresos, 'filas_gastos': filas_gastos, 'filas_balance': filas_balance,
        'total_ingreso_presupuestado': sum((f['presupuestado'] for f in filas_ingresos), Decimal('0.00')),
        'total_ingreso_recaudado': sum((f['recaudado'] for f in filas_ingresos), Decimal('0.00')),
        'total_gasto_apropiado': sum((f['apropiacion_definitiva'] for f in filas_gastos), Decimal('0.00')),
        'total_gasto_comprometido': sum((f['comprometido'] for f in filas_gastos), Decimal('0.00')),
        'total_gasto_obligado': sum((f['obligado'] for f in filas_gastos), Decimal('0.00')),
        'total_gasto_pagado': sum((f['pagado'] for f in filas_gastos), Decimal('0.00')),
    })


@login_required
def exportar_reporte_ejecucion_excel(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    filtro = _filtro_institucion(request)
    vigencia, _vigencias = _resolver_vigencia(request, filtro)
    if not vigencia:
        messages.error(request, 'No hay ninguna vigencia fiscal para exportar.')
        return redirect('presupuesto:reporte_ejecucion')

    def _hoja(wb, titulo, encabezados, filas):
        ws = wb.create_sheet(titulo)
        ws.append(encabezados)
        for cell in ws[1]:
            cell.fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF", size=10)
            cell.alignment = Alignment(horizontal='center')
        for fila in filas:
            ws.append(fila)
        for col_cells in ws.columns:
            w = max((len(str(c.value or '')) for c in col_cells), default=8)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max(w + 2, 10), 40)
        ws.freeze_panes = "A2"
        return ws

    wb = Workbook()
    wb.remove(wb.active)

    _hoja(wb, 'Ejecución Ingresos', ['Rubro', 'Fuente de Financiación (CHIP)', 'Presupuestado', 'Recaudado', 'Saldo por Recaudar'], [
        [str(f['rubro']), str(f['fuente_financiacion'] or ''), float(f['presupuestado']), float(f['recaudado']), float(f['saldo_por_recaudar'])]
        for f in reportes.ejecucion_ingresos(vigencia)
    ])
    _hoja(wb, 'Ejecución Gastos', ['Rubro', 'Apropiación Inicial', 'Apropiación Definitiva', 'Comprometido (RP)', 'Obligado', 'Pagado', 'Saldo por Comprometer'], [
        [str(f['rubro']), float(f['apropiacion_inicial']), float(f['apropiacion_definitiva']), float(f['comprometido']), float(f['obligado']), float(f['pagado']), float(f['saldo_por_comprometer'])]
        for f in reportes.ejecucion_gastos(vigencia)
    ])
    _hoja(wb, 'Balance Comprobación', ['Cuenta CGC', 'Nombre', 'Total Débitos', 'Total Créditos', 'Saldo'], [
        [f['cuenta__codigo'], f['cuenta__nombre'], float(f['total_debitos']), float(f['total_creditos']), float(f['saldo'])]
        for f in reportes.balance_comprobacion(vigencia)
    ])

    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="reporte_ejecucion_presupuestal_{vigencia.anio}.xlsx"'
    wb.save(resp)
    return resp


@login_required
def exportar_reporte_ejecucion_pdf(request):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    filtro = _filtro_institucion(request)
    vigencia, _vigencias = _resolver_vigencia(request, filtro)
    if not vigencia:
        messages.error(request, 'No hay ninguna vigencia fiscal para exportar.')
        return redirect('presupuesto:reporte_ejecucion')

    from django.template.loader import get_template
    from django.utils import timezone
    from xhtml2pdf import pisa

    from .pdf_utils import link_callback_pdf

    filas_ingresos = reportes.ejecucion_ingresos(vigencia)
    filas_gastos = reportes.ejecucion_gastos(vigencia)
    filas_balance = reportes.balance_comprobacion(vigencia)
    context = {
        'institucion': _get_institucion(request),
        'vigencia': vigencia,
        'filas_ingresos': filas_ingresos,
        'filas_gastos': filas_gastos,
        'filas_balance': filas_balance,
        'total_ingreso_presupuestado': sum((f['presupuestado'] for f in filas_ingresos), Decimal('0.00')),
        'total_ingreso_recaudado': sum((f['recaudado'] for f in filas_ingresos), Decimal('0.00')),
        'total_gasto_apropiado': sum((f['apropiacion_definitiva'] for f in filas_gastos), Decimal('0.00')),
        'total_gasto_comprometido': sum((f['comprometido'] for f in filas_gastos), Decimal('0.00')),
        'total_gasto_obligado': sum((f['obligado'] for f in filas_gastos), Decimal('0.00')),
        'total_gasto_pagado': sum((f['pagado'] for f in filas_gastos), Decimal('0.00')),
        'fecha_generacion': timezone.now(),
    }
    html = get_template('presupuesto/pdfs/reporte_ejecucion.html').render(context)
    resp = HttpResponse(content_type='application/pdf')
    resp['Content-Disposition'] = f'inline; filename="reporte_ejecucion_presupuestal_{vigencia.anio}.pdf"'
    pisa_status = pisa.CreatePDF(html, dest=resp, link_callback=link_callback_pdf)
    if pisa_status.err:
        return HttpResponse('Ocurrió un error al generar el reporte en PDF.', status=500)
    return resp


# ─────────────────────────────────────────────────────────────────────────
# Documentos imprimibles (PDF) — CDP, RP, Obligación, Orden de Pago,
# Comprobante de Egreso. Cada uno es el soporte físico/legal del expediente
# contable del Fondo de Servicios Educativos.
# ─────────────────────────────────────────────────────────────────────────

def _renderizar_pdf(request, template_name, context, nombre_archivo):
    from django.template.loader import get_template
    from django.utils import timezone
    from xhtml2pdf import pisa

    from .pdf_utils import link_callback_pdf

    context = {**context, 'institucion': _get_institucion(request), 'fecha_generacion': timezone.now()}
    html = get_template(template_name).render(context)
    resp = HttpResponse(content_type='application/pdf')
    resp['Content-Disposition'] = f'inline; filename="{nombre_archivo}"'
    resp['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
    pisa_status = pisa.CreatePDF(html, dest=resp, link_callback=link_callback_pdf)
    if pisa_status.err:
        return HttpResponse('Ocurrió un error al generar el PDF.', status=500)
    return resp


@login_required
def imprimir_cdp(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    from .pdf_utils import valor_en_letras
    cdp = get_object_or_404(CDP.objects.select_related('apropiacion__rubro', 'vigencia', 'creado_por'), pk=pk, **_filtro_institucion(request))
    return _renderizar_pdf(request, 'presupuesto/pdfs/cdp.html', {
        'cdp': cdp, 'valor_letras': valor_en_letras(cdp.valor),
    }, f'CDP_{cdp.numero}_{cdp.vigencia.anio}.pdf')


@login_required
def imprimir_rp(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    from .pdf_utils import valor_en_letras
    rp = get_object_or_404(
        RP.objects.select_related('cdp__vigencia', 'tercero', 'categoria_cpc', 'creado_por'),
        pk=pk, **_filtro_institucion(request),
    )
    return _renderizar_pdf(request, 'presupuesto/pdfs/rp.html', {
        'rp': rp, 'valor_letras': valor_en_letras(rp.valor),
    }, f'RP_{rp.numero}.pdf')


@login_required
def imprimir_obligacion(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    from .pdf_utils import valor_en_letras
    obligacion = get_object_or_404(
        Obligacion.objects.select_related('rp__tercero', 'creado_por'),
        pk=pk, **_filtro_institucion(request),
    )
    return _renderizar_pdf(request, 'presupuesto/pdfs/obligacion.html', {
        'obligacion': obligacion, 'valor_letras': valor_en_letras(obligacion.valor),
    }, f'Obligacion_{obligacion.numero}.pdf')


@login_required
def imprimir_orden_pago(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    from .pdf_utils import valor_en_letras
    orden = get_object_or_404(
        OrdenDePago.objects.select_related('obligacion__rp', 'beneficiario', 'creado_por').prefetch_related('retenciones__concepto'),
        pk=pk, **_filtro_institucion(request),
    )
    return _renderizar_pdf(request, 'presupuesto/pdfs/orden_pago.html', {
        'orden': orden, 'valor_letras': valor_en_letras(orden.valor_neto),
    }, f'OrdenDePago_{orden.numero}.pdf')


@login_required
def imprimir_comprobante(request, pk):
    guard = _requiere_gestor(request)
    if guard:
        return guard
    comprobante = get_object_or_404(
        ComprobanteContable.objects.select_related(
            'vigencia', 'orden_pago__beneficiario', 'comprobante_que_reversa', 'cuenta_bancaria', 'contabilizado_por',
        ).prefetch_related('movimientos__cuenta'),
        pk=pk, **_filtro_institucion(request),
    )
    return _renderizar_pdf(request, 'presupuesto/pdfs/comprobante.html', {
        'comprobante': comprobante,
    }, f'Comprobante_{comprobante.numero}.pdf')
